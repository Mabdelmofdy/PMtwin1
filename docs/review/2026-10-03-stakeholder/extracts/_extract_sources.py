"""Read-only extraction of stakeholder source documents. Does not modify sources."""
from __future__ import annotations

import json
from pathlib import Path

import fitz
import openpyxl
from docx import Document

OUT = Path(__file__).resolve().parent
DOWNLOADS = Path(r"C:\Users\Mohammed Saleh\Downloads")

PDF = DOWNLOADS / "PMTwin-Need&Offer Framework-Matching Mechanizim-10Jan2025.pdf"
HANDOVER = DOWNLOADS / "PM-Twin_Project_Master_Context_Handover.docx"
CR_REVIEW = DOWNLOADS / "PM-Twin_Wireframes_Core_CR_Review.docx"
COMPANY = DOWNLOADS / "PMTwin Matching Wireframes-10Aug2026 - Companyflow.xlsx"
INDIVIDUAL = DOWNLOADS / "PMTwin Matching Wireframes-10Aug2026-Indivadual.xlsx"


def extract_pdf(path: Path, dest: Path) -> None:
    doc = fitz.open(path)
    parts = [f"# {path.name}\n", f"pages: {doc.page_count}\n"]
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        parts.append(f"\n\n===== PAGE {i} =====\n")
        parts.append(text)
    dest.write_text("".join(parts), encoding="utf-8")
    print(f"pdf pages={doc.page_count} chars={dest.stat().st_size}")
    doc.close()


def extract_docx(path: Path, dest: Path) -> None:
    doc = Document(path)
    lines = [f"# {path.name}", ""]
    for i, p in enumerate(doc.paragraphs, start=1):
        style = p.style.name if p.style is not None else ""
        text = p.text.replace("\u00a0", " ")
        if text.strip() or style.startswith("Heading"):
            lines.append(f"[p{i}|{style}] {text}")
    lines.append("\n## TABLES\n")
    for ti, table in enumerate(doc.tables, start=1):
        lines.append(f"\n### TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}\n")
        for ri, row in enumerate(table.rows, start=1):
            cells = [c.text.replace("\n", " | ").strip() for c in row.cells]
            lines.append(f"R{ri}: " + " || ".join(cells))
    dest.write_text("\n".join(lines), encoding="utf-8")
    print(f"docx {path.name} paras={len(doc.paragraphs)} tables={len(doc.tables)} bytes={dest.stat().st_size}")


def cell_addr(cell) -> str:
    return cell.coordinate


def extract_xlsx(path: Path, dest_dir: Path) -> None:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    summary = {
        "file": path.name,
        "sheets": [],
        "defined_names": [],
    }
    try:
        for dn in wb.defined_names.definedName:
            summary["defined_names"].append({"name": dn.name, "attr": dn.attr_text})
    except Exception as exc:
        summary["defined_names_error"] = str(exc)

    for ws in wb.worksheets:
        info = {
            "title": ws.title,
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sheet_state": ws.sheet_state,
            "merged": [str(m) for m in ws.merged_cells.ranges],
        }
        summary["sheets"].append(info)
        lines = [
            f"# {path.name} :: {ws.title}",
            f"state={ws.sheet_state} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}",
            f"merged={info['merged'][:80]}",
            "",
        ]
        # data validations
        dvs = getattr(ws, "data_validations", None)
        if dvs is not None:
            lines.append("## DATA VALIDATIONS")
            for dv in dvs.dataValidation:
                lines.append(
                    f"sqref={dv.sqref} type={dv.type} formula1={dv.formula1} formula2={dv.formula2} allow_blank={dv.allow_blank}"
                )
            lines.append("")
        lines.append("## CELLS")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=ws.max_column or 1):
            for cell in row:
                if cell.value is None and (cell.data_type is None or cell.data_type == "n"):
                    continue
                val = cell.value
                if isinstance(val, str):
                    val = val.replace("\r\n", "\n").replace("\n", " \\n ")
                lines.append(f"{cell.coordinate}\t{cell.data_type}\t{val}")
        safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in ws.title)[:80]
        out = dest_dir / f"{path.stem[:40]}__{safe}.txt"
        out.write_text("\n".join(lines), encoding="utf-8")
        print(f"sheet {ws.title} rows={ws.max_row} cols={ws.max_column} bytes={out.stat().st_size}")
    (dest_dir / f"{path.stem[:40]}__INDEX.json").write_text(
        json.dumps(summary, indent=2, default=str), encoding="utf-8"
    )
    wb.close()


def main() -> None:
    extract_pdf(PDF, OUT / "need-offer-framework-2025-01-10.txt")
    extract_docx(HANDOVER, OUT / "handover.txt")
    extract_docx(CR_REVIEW, OUT / "wireframes-core-cr-review.txt")
    company_dir = OUT / "company"
    indiv_dir = OUT / "individual"
    company_dir.mkdir(exist_ok=True)
    indiv_dir.mkdir(exist_ok=True)
    extract_xlsx(COMPANY, company_dir)
    extract_xlsx(INDIVIDUAL, indiv_dir)


if __name__ == "__main__":
    main()
